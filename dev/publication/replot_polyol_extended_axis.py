from __future__ import annotations

import argparse
import json
import os
import time
from dataclasses import dataclass
from pathlib import Path
from typing import Any

import numpy as np
import pandas as pd
from datapackage import Package
from openpyxl import load_workbook

from trails import Trails, get_lcia_method_names

PUBLICATION_DIR = Path(__file__).resolve().parent
DATAPACKAGE = PUBLICATION_DIR / "trails_remind_SSP2-PkBudg1000.zip"
LCI_DIR = PUBLICATION_DIR / "LCIs"
OUTPUT_DIR = PUBLICATION_DIR / "notebook_runs" / "temporal_lci_routing_comparison_clean"

REFERENCE_YEAR = 2025
EI_VERSION = "3.12"
METHOD_PREFIX = "EF v3.1 - "
TARGET_METHOD_BY_CASE = {
    "bev": (
        "EF v3.1 - ecotoxicity: freshwater, inorganics - "
        "comparative toxic unit for ecosystems (CTUe)"
    ),
    "polyol": (
        "EF v3.1 - material resources: metals/minerals - "
        "abiotic depletion potential (ADP): elements (ultimate reserves)"
    ),
}

DEMAND_AMOUNT_BY_CASE = {
    "bev": 150_000.0,
    "polyol": 50_000_000_000.0,
}

FOREGROUND_DEPTH = 1
ADAPTIVE_RELATIVE_SCORE_CUTOFF = 1e-4
ADAPTIVE_MIN_DEPTH = 1
ADAPTIVE_MAX_DEPTH = None
ROUTING_MIN_AMOUNT = 0.0
SOLVER_MODE = "iterative"
FALLBACK_SOLVER_MODE = "direct"
SHOW_PROGRESS = False

INTERPOLATION_START_YEAR_OFFSET = -20
INTERPOLATION_END_YEAR_OFFSET = 20

PLOT_TOP_ACTIVITIES = 8
PLOT_YEAR_START_FALLBACK = 1985
X_TICK_START = 1990
X_TICK_YEARS = 10
WIDTH = 640
HEIGHT = 504
PNG_SCALE = 3
LEGEND_COLUMNS = 2
LEGEND_MAX_ROWS = 5
LEGEND_LABEL_MAX_CHARS = 50

ROUTING_CASE_LABELS = {
    "foreground": "Foreground only",
    "adaptive": "Adaptive routing",
}
ROUTING_CASE_COLORS = {
    "foreground": "#FFA500",
    "adaptive": "#FF0000",
    "static": "#dc2626",
}

EXPECTED_INVENTORY_ACTIVITIES = {
    "lci-case-study-ccu_polyol_delayed_release.xlsx": {
        "polyol precursor production from captured CO2",
        "treatment of polyol, incineration",
    },
    "lci-case-study-daccs_storage_risk.xlsx": {
        "carbon dioxide, captured, with a solvent-based direct air capture system, 1MtCO2",
    },
    "lci-case-study-marine_fuel_switch.xlsx": {
        "marine freight service, temporal fuel transition",
    },
    "lci-pass_cars.xlsx": {
        "transport, passenger, car, diesel",
        "transport, passenger, car, battery electric",
        "heating, from wood logs stove",
        "heating, from heat pump",
    },
}


@dataclass(frozen=True)
class ActivityDef:
    name: str
    reference_product: str
    location: str


@dataclass
class RoutingResult:
    routing: str
    activity_index: int
    activity_label: str
    method: str
    static_score: float
    scores: Any
    annual: pd.Series
    cumulative: pd.Series
    routing_seconds: float
    lca_seconds: float
    solver_used: str
    graph_nodes: int
    graph_edges: int
    deepest_level: int


ACTIVITY_BY_CASE = {
    "bev": ActivityDef(
        "transport, passenger, car, battery electric",
        "transport, passenger, car",
        "RER",
    ),
    "polyol": ActivityDef(
        "polyol precursor production from captured CO2",
        "polyol precursor",
        "RER",
    ),
}


def clean_text(value: object) -> str:
    return "" if value is None else str(value).strip()


def inventory_activity_names(path: Path) -> set[str]:
    workbook = load_workbook(path, data_only=True, read_only=True)
    names: set[str] = set()
    for worksheet in workbook.worksheets:
        for row in worksheet.iter_rows(values_only=True):
            values = [cell for cell in row if cell is not None]
            if len(values) >= 2 and str(values[0]).strip() == "Activity":
                names.add(str(values[1]).strip())
    return names


def validate_configuration(inventory_paths: list[Path]) -> None:
    missing_paths = [
        path for path in [DATAPACKAGE, *inventory_paths] if not path.exists()
    ]
    if missing_paths:
        raise FileNotFoundError(
            "Missing input files:\n" + "\n".join(map(str, missing_paths))
        )

    paths_by_name = {path.name: path for path in inventory_paths}
    missing_files = sorted(set(EXPECTED_INVENTORY_ACTIVITIES) - set(paths_by_name))
    if missing_files:
        raise FileNotFoundError(
            "Missing publication inventory files:\n" + "\n".join(missing_files)
        )

    problems: list[str] = []
    for filename, expected_names in EXPECTED_INVENTORY_ACTIVITIES.items():
        actual_names = inventory_activity_names(paths_by_name[filename])
        missing_names = sorted(expected_names - actual_names)
        if missing_names:
            problems.append(
                f"{filename} is not the expected Figure 5 input; "
                f"missing activities: {missing_names}"
            )
    if problems:
        raise ValueError("\n".join(problems))


def get_available_methods(trails: Trails) -> list[str]:
    if hasattr(trails, "get_available_methods"):
        try:
            return list(trails.get_available_methods(ei_version=EI_VERSION))
        except TypeError:
            return list(trails.get_available_methods())
    if hasattr(trails, "list_lcia_methods"):
        return list(trails.list_lcia_methods(ei_version=EI_VERSION))
    return list(get_lcia_method_names(ei_version=EI_VERSION))


def get_ef_v31_methods(trails: Trails) -> list[str]:
    methods = [
        method
        for method in get_available_methods(trails)
        if str(method).startswith(METHOD_PREFIX)
    ]
    if not methods:
        raise ValueError(f"No available methods start with {METHOD_PREFIX!r}")
    return sorted(methods)


def matches_activity(meta: dict[str, Any], target: ActivityDef) -> bool:
    return (
        clean_text(meta.get("name")) == target.name
        and clean_text(meta.get("reference product")) == target.reference_product
        and clean_text(meta.get("location")) == target.location
    )


def find_activity_index(trails: Trails, target: ActivityDef) -> int:
    for metadata_by_index in trails.activity_indices.values():
        for index, meta in metadata_by_index.items():
            if matches_activity(meta, target):
                return int(index)
    raise ValueError(
        "Could not find activity: "
        f"{target.name} | {target.reference_product} | {target.location}"
    )


def activity_label(trails: Trails, activity_index: int) -> str:
    for metadata_by_index in trails.activity_indices.values():
        meta = metadata_by_index.get(int(activity_index))
        if meta:
            return " | ".join(
                part
                for part in [
                    clean_text(meta.get("name")),
                    clean_text(meta.get("reference product")),
                    clean_text(meta.get("location")),
                ]
                if part
            )
    return f"activity {activity_index}"


def graph_size(trails: Trails) -> tuple[int, int]:
    graph = getattr(trails, "graph", None)
    if graph is None:
        return 0, 0
    return len(graph.nodes), len(graph.edges)


def deepest_graph_level(trails: Trails) -> int:
    graph = getattr(trails, "graph", None)
    if graph is None:
        return 0
    depths = [int(data.get("depth", 0)) for _node, data in graph.nodes(data=True)]
    return max(depths, default=0)


def dense_1d(data: Any) -> np.ndarray:
    raw = data.data if hasattr(data, "data") else data
    if hasattr(raw, "todense"):
        return np.asarray(raw.todense(), dtype=float).ravel()
    return np.asarray(data, dtype=float).ravel()


def select_method(scores: Any, method: str) -> Any:
    if "method" not in scores.dims:
        return scores
    methods = [str(value) for value in scores.coords["method"].values.tolist()]
    return scores.isel(method=methods.index(method), drop=True)


def reduce_scores(scores: Any, keep_dims: tuple[str, ...]) -> Any:
    kept = [dim for dim in keep_dims if dim in scores.dims]
    reduce_dims = [dim for dim in scores.dims if dim not in kept]
    reduced = scores.sum(dim=reduce_dims) if reduce_dims else scores
    return reduced.transpose(*kept) if tuple(reduced.dims) != tuple(kept) else reduced


def annual_total_series(scores: Any, method: str) -> pd.Series:
    by_year = reduce_scores(select_method(scores, method), ("year",))
    years = np.asarray(by_year.coords["year"].values, dtype=int)
    return pd.Series(dense_1d(by_year), index=years).sort_index()


def root_activity_frame(scores: Any, method: str) -> pd.DataFrame:
    by_root_year = reduce_scores(
        select_method(scores, method), ("root activity", "year")
    )
    root_ids = np.asarray(by_root_year.coords["root activity"].values, dtype=int)
    years = np.asarray(by_root_year.coords["year"].values, dtype=int)
    values = np.asarray(
        (
            by_root_year.data.todense()
            if hasattr(by_root_year.data, "todense")
            else by_root_year.values
        ),
        dtype=float,
    )
    return pd.DataFrame(values, index=root_ids, columns=years)


def static_score_as_float(score: object) -> float:
    values = np.asarray(score, dtype=float).ravel()
    if values.size == 0:
        raise ValueError("Static score is empty")
    return float(values[0])


def slug(value: str, max_length: int = 120) -> str:
    text = "".join(char.lower() if char.isalnum() else "_" for char in str(value))
    text = "_".join(part for part in text.split("_") if part)
    return (text or "item")[:max_length]


def run_temporal_case(
    *,
    trails: Trails,
    routing: str,
    activity_index: int,
    activity_label_: str,
    amount: float,
    method: str,
    static_score: float,
    adaptive_methods: list[str],
) -> RoutingResult:
    if routing == "foreground_depth_1":
        routing_kwargs = dict(
            max_depth=FOREGROUND_DEPTH,
            adaptive_relative_score_cutoff=None,
        )
    elif routing == "adaptive":
        routing_kwargs = dict(
            max_depth=ADAPTIVE_MAX_DEPTH,
            adaptive_relative_score_cutoff=ADAPTIVE_RELATIVE_SCORE_CUTOFF,
            adaptive_methods=adaptive_methods,
            adaptive_ei_version=EI_VERSION,
            adaptive_min_depth=ADAPTIVE_MIN_DEPTH,
        )
    else:
        raise ValueError(f"Unknown routing case: {routing}")

    t0 = time.perf_counter()
    trails.temporal_routing(
        start_year=REFERENCE_YEAR,
        start_act_idx=activity_index,
        amount=amount,
        min_amount=ROUTING_MIN_AMOUNT,
        show_progress=SHOW_PROGRESS,
        attribute_to_roots=True,
        **routing_kwargs,
    )
    routing_seconds = time.perf_counter() - t0
    nodes, edges = graph_size(trails)
    deepest_level = deepest_graph_level(trails)

    solver_used = SOLVER_MODE
    t0 = time.perf_counter()
    try:
        trails.lca(
            methods=[method],
            ei_version=EI_VERSION,
            show_progress=SHOW_PROGRESS,
            attribute_to_roots=True,
            compute_score=True,
            store_inventory=False,
            solver_mode=SOLVER_MODE,
        )
    except RuntimeError as error:
        can_retry_direct = (
            "GMRES failed to converge" in str(error)
            and FALLBACK_SOLVER_MODE == "direct"
            and SOLVER_MODE != "direct"
        )
        if not can_retry_direct:
            raise
        solver_used = "direct"
        trails.lca(
            methods=[method],
            ei_version=EI_VERSION,
            show_progress=SHOW_PROGRESS,
            attribute_to_roots=True,
            compute_score=True,
            store_inventory=False,
            solver_mode="direct",
        )
    lca_seconds = time.perf_counter() - t0

    scores = reduce_scores(trails.scores, ("method", "root activity", "year")).copy(
        deep=True
    )
    annual = annual_total_series(scores, method)
    return RoutingResult(
        routing=routing,
        activity_index=activity_index,
        activity_label=activity_label_,
        method=method,
        static_score=static_score,
        scores=scores,
        annual=annual,
        cumulative=annual.cumsum(),
        routing_seconds=routing_seconds,
        lca_seconds=lca_seconds,
        solver_used=solver_used,
        graph_nodes=nodes,
        graph_edges=edges,
        deepest_level=deepest_level,
    )


def top_root_contributions(result: RoutingResult) -> pd.DataFrame:
    frame = root_activity_frame(result.scores, result.method)
    totals = frame.abs().sum(axis=1).sort_values(ascending=False)
    top_roots = list(totals.head(PLOT_TOP_ACTIVITIES).index)
    plotted = frame.loc[top_roots].copy()
    remaining = frame.drop(index=top_roots, errors="ignore")
    if not remaining.empty:
        plotted.loc[-1] = remaining.sum(axis=0)
    return plotted


def method_unit(method: str) -> str:
    from trails.lcia import _get_lcia_methods_filepath

    filepath = _get_lcia_methods_filepath(EI_VERSION)
    with open(filepath) as handle:
        for row in json.load(handle):
            if " - ".join(row.get("name", [])) == method:
                return clean_text(row.get("unit")) or "impact units"
    if "(" in method and method.endswith(")"):
        return method.rsplit("(", 1)[-1].removesuffix(")")
    return "impact units"


def short_label(text: str) -> str:
    text = str(text)
    limit = int(LEGEND_LABEL_MAX_CHARS)
    if len(text) <= limit:
        return text
    return text[: max(0, limit - 3)].rstrip() + "..."


def reference_product_label(trails: Trails, activity_index: int) -> str:
    for metadata_by_index in trails.activity_indices.values():
        meta = metadata_by_index.get(int(activity_index))
        if meta:
            product = clean_text(meta.get("reference product"))
            return product or f"activity {activity_index}"
    return f"activity {activity_index}"


def contribution_palette() -> list[str]:
    return [
        "#636EFA",
        "#EF553B",
        "#00CC96",
        "#AB63FA",
        "#FFA15A",
        "#19D3F3",
        "#FF6692",
        "#B6E880",
        "#FF97FF",
        "#FECB52",
    ]


def routing_kind(routing: str) -> str:
    return "foreground" if routing.startswith("foreground") else routing


def rgba_tuple(color: str, alpha: float) -> tuple[float, float, float, float]:
    text = str(color).lstrip("#")
    return (
        int(text[0:2], 16) / 255,
        int(text[2:4], 16) / 255,
        int(text[4:6], 16) / 255,
        alpha,
    )


def positive_y_range(values: list[np.ndarray | pd.Series]) -> list[float]:
    arrays = [np.asarray(value, dtype=float).ravel() for value in values]
    finite_arrays = [array[np.isfinite(array)] for array in arrays if array.size]
    if not finite_arrays:
        return [0.0, 1.0]
    finite_values = np.concatenate(finite_arrays)
    if finite_values.size == 0:
        return [0.0, 1.0]
    y_max = max(0.0, float(finite_values.max()))
    if y_max <= 0:
        return [0.0, 1.0]
    return [0.0, y_max * 1.1]


def nonzero_years(values: pd.Series | pd.DataFrame) -> list[int]:
    if isinstance(values, pd.Series):
        mask = values.abs() > 1e-12
        return [int(year) for year in values.index[mask].tolist()]
    mask = values.abs().gt(1e-12).any(axis=0)
    return [int(year) for year in values.columns[mask].tolist()]


def impact_year_bounds(
    *series: pd.Series,
    contributions: pd.DataFrame,
) -> tuple[int, int]:
    years: list[int] = []
    for item in series:
        years.extend(nonzero_years(item))
    years.extend(nonzero_years(contributions))
    if not years:
        return PLOT_YEAR_START_FALLBACK, PLOT_YEAR_START_FALLBACK
    return min(years), max(years)


def write_series_csv(
    *,
    foreground: RoutingResult,
    adaptive: RoutingResult,
    output_path: Path,
) -> None:
    all_years = sorted(set(foreground.annual.index).union(set(adaptive.annual.index)))
    annual_total = (
        foreground.annual.reindex(all_years, fill_value=0.0).abs()
        + adaptive.annual.reindex(all_years, fill_value=0.0).abs()
    )
    nonzero = annual_total[annual_total > 1e-12]
    if nonzero.empty:
        years = all_years
    else:
        start = int(nonzero.index.min())
        end = int(nonzero.index.max())
        years = list(range(start, end + 1))
    frame = pd.DataFrame(index=years)
    frame.index.name = "year"
    frame["foreground_annual"] = foreground.annual.reindex(years, fill_value=0.0)
    frame["foreground_cumulative"] = foreground.cumulative.reindex(years).ffill()
    frame["adaptive_annual"] = adaptive.annual.reindex(years, fill_value=0.0)
    frame["adaptive_cumulative"] = adaptive.cumulative.reindex(years).ffill()
    frame.to_csv(output_path)


def case_output_stem(method: str) -> str:
    return f"{slug(method, 100)}_routing_comparison"


def write_case_png(
    *,
    trails: Trails,
    foreground: RoutingResult,
    adaptive: RoutingResult,
    stacked: bool,
    path: Path,
) -> tuple[int, int]:
    import matplotlib

    os.environ.setdefault("MPLCONFIGDIR", str(OUTPUT_DIR / ".matplotlib"))
    Path(os.environ["MPLCONFIGDIR"]).mkdir(parents=True, exist_ok=True)

    matplotlib.use("Agg", force=True)
    import matplotlib.pyplot as plt
    from matplotlib.lines import Line2D
    from matplotlib.patches import Patch

    contributions = top_root_contributions(adaptive)
    palette = contribution_palette()
    unit = method_unit(adaptive.method)
    plot_year_start, plot_year_end = impact_year_bounds(
        foreground.annual,
        adaptive.annual,
        contributions=contributions,
    )

    figure, axis = plt.subplots(figsize=(WIDTH / 100, HEIGHT / 100), dpi=100)
    figure.patch.set_facecolor("white")
    cumulative_axis = axis.twinx()
    legend_handles = []
    legend_labels = []

    years = np.asarray(contributions.columns, dtype=float)
    baseline = np.zeros_like(years, dtype=float)
    for position, (root_id, values) in enumerate(contributions.iterrows()):
        color = palette[position % len(palette)]
        label = (
            "Other activities"
            if int(root_id) == -1
            else reference_product_label(trails, int(root_id))
        )
        annual = values.to_numpy(dtype=float)
        if stacked:
            next_baseline = baseline + annual
            axis.fill_between(
                years,
                baseline,
                next_baseline,
                color=rgba_tuple(color, 0.38),
                linewidth=0,
            )
            baseline = next_baseline
        else:
            axis.fill_between(
                years, 0, annual, color=rgba_tuple(color, 0.38), linewidth=0
            )
        legend_handles.append(
            Patch(facecolor=rgba_tuple(color, 0.75), edgecolor="none")
        )
        legend_labels.append(short_label(label))

    for result in [foreground, adaptive]:
        kind = routing_kind(result.routing)
        color = ROUTING_CASE_COLORS[kind]
        label = ROUTING_CASE_LABELS[kind]
        line_width = 3.8 if kind == "adaptive" else 3.2
        cumulative_axis.plot(
            result.cumulative.index,
            result.cumulative.values,
            color=color,
            linewidth=line_width,
        )
        legend_handles.append(Line2D([0], [0], color=color, linewidth=line_width))
        legend_labels.append(short_label(f"{label} (cumul.)"))

    cumulative_axis.axhline(
        adaptive.static_score,
        color=ROUTING_CASE_COLORS["static"],
        linewidth=2.4,
        linestyle="--",
    )
    legend_handles.append(
        Line2D(
            [0],
            [0],
            color=ROUTING_CASE_COLORS["static"],
            linewidth=2.4,
            linestyle="--",
        )
    )
    legend_labels.append(short_label("Static score (2025)"))
    axis.axvline(
        REFERENCE_YEAR,
        color="black",
        alpha=0.45,
        linewidth=1.0,
        linestyle="--",
        zorder=4,
    )
    axis.annotate(
        "reference year",
        xy=(REFERENCE_YEAR, 1.0),
        xycoords=("data", "axes fraction"),
        xytext=(3, -4),
        textcoords="offset points",
        ha="left",
        va="top",
        fontsize=7,
        color="black",
        alpha=0.65,
    )

    axis.set_xlabel("Year", fontsize=10, color="#1f3557")
    axis.set_ylabel(f"Annual impact ({unit})", fontsize=10, color="#1f3557")
    cumulative_axis.set_ylabel(
        f"Cumulative / static impact ({unit})",
        fontsize=10,
        color="#1f3557",
        labelpad=8,
    )
    tick_start = int(np.ceil(plot_year_start / X_TICK_YEARS) * X_TICK_YEARS)
    tick_start = max(tick_start, X_TICK_START)
    axis.set_xticks(np.arange(tick_start, plot_year_end + 1, X_TICK_YEARS))
    axis.set_xlim(plot_year_start, plot_year_end)
    primary_y_values = [baseline] if stacked else [contributions.to_numpy(dtype=float)]
    axis.set_ylim(positive_y_range(primary_y_values))
    cumulative_axis.set_ylim(
        positive_y_range(
            [
                foreground.cumulative,
                adaptive.cumulative,
                np.array([adaptive.static_score], dtype=float),
            ]
        )
    )
    axis.set_axisbelow(True)
    axis.grid(True, color="#e6eef8")
    cumulative_axis.grid(False)
    for chart_axis in [axis, cumulative_axis]:
        chart_axis.tick_params(axis="y", labelsize=10, colors="#1f3557")
        chart_axis.tick_params(axis="x", labelsize=8, colors="#1f3557")
        chart_axis.spines["top"].set_visible(False)

    figure.subplots_adjust(left=0.13, right=0.80, top=0.66, bottom=0.14)
    max_entries = int(LEGEND_COLUMNS) * int(LEGEND_MAX_ROWS)
    if len(legend_handles) > max_entries and len(legend_handles) > 3:
        visible_handles = legend_handles[: max_entries - 3] + legend_handles[-3:]
        visible_labels = legend_labels[: max_entries - 3] + legend_labels[-3:]
    else:
        visible_handles = legend_handles[:max_entries]
        visible_labels = legend_labels[:max_entries]
    figure.legend(
        visible_handles,
        visible_labels,
        loc="lower left",
        bbox_to_anchor=(0.08, 0.72, 0.86, 0.26),
        ncol=2,
        frameon=False,
        fontsize=7.5,
        handlelength=1.6,
        columnspacing=0.6,
        handletextpad=0.5,
        labelcolor="#1f3557",
    )

    path.parent.mkdir(parents=True, exist_ok=True)
    figure.savefig(path, dpi=100 * PNG_SCALE, facecolor="white")
    plt.close(figure)
    return plot_year_start, plot_year_end


def parse_args() -> argparse.Namespace:
    parser = argparse.ArgumentParser(
        description=(
            "Regenerate selected Figure 5 routing-comparison panels with the "
            "time axis set from first to last nonzero annual impact."
        )
    )
    parser.add_argument(
        "--case",
        choices=sorted(TARGET_METHOD_BY_CASE),
        action="append",
        dest="cases",
        help="Case key to regenerate. Can be provided more than once.",
    )
    return parser.parse_args()


def main() -> None:
    args = parse_args()
    target_cases = args.cases or ["polyol"]

    inventory_paths = sorted(path for path in LCI_DIR.glob("lci-*.xlsx"))
    validate_configuration(inventory_paths)
    OUTPUT_DIR.mkdir(parents=True, exist_ok=True)

    print(f"Loading datapackage: {DATAPACKAGE}")
    trails = Trails(
        Package(str(DATAPACKAGE)),
        interpolate_annual=True,
        cache_interpolation=True,
        interpolation_start_year_offset=INTERPOLATION_START_YEAR_OFFSET,
        interpolation_end_year_offset=INTERPOLATION_END_YEAR_OFFSET,
        ei_version=EI_VERSION,
    )
    ef_methods = get_ef_v31_methods(trails)
    for case_key in target_cases:
        method = TARGET_METHOD_BY_CASE[case_key]
        if method not in ef_methods:
            raise ValueError(f"Target method is unavailable: {method}")
    print(f"Selected {len(ef_methods)} EF v3.1 methods for adaptive routing")

    print("Importing foreground inventories")
    trails.import_excel_inventory([str(path) for path in inventory_paths])

    for case_key in target_cases:
        method = TARGET_METHOD_BY_CASE[case_key]
        activity_index = find_activity_index(trails, ACTIVITY_BY_CASE[case_key])
        label = activity_label(trails, activity_index)
        amount = DEMAND_AMOUNT_BY_CASE[case_key]
        print(f"Case: {case_key}")
        print(f"Activity: {label}")
        print(f"Method: {method}")
        print(f"Amount: {amount:g}")

        t0 = time.perf_counter()
        trails.static_lca(
            year=REFERENCE_YEAR,
            act_idx=activity_index,
            methods=[method],
            amount=amount,
            ei_version=EI_VERSION,
        )
        static_score = static_score_as_float(trails.static_score)
        print(f"Static score: {static_score:.12g} ({time.perf_counter() - t0:.1f}s)")

        results = {}
        for routing in ("foreground_depth_1", "adaptive"):
            result = run_temporal_case(
                trails=trails,
                routing=routing,
                activity_index=activity_index,
                activity_label_=label,
                amount=amount,
                method=method,
                static_score=static_score,
                adaptive_methods=ef_methods,
            )
            results[routing] = result
            print(
                f"{routing}: cumulative={result.cumulative.iloc[-1]:.12g}, "
                f"depth={result.deepest_level}, nodes={result.graph_nodes}, "
                f"edges={result.graph_edges}, solver={result.solver_used}, "
                f"routing={result.routing_seconds:.1f}s, lca={result.lca_seconds:.1f}s"
            )

        foreground = results["foreground_depth_1"]
        adaptive = results["adaptive"]

        figure_dir = OUTPUT_DIR / case_key
        stem = case_output_stem(method)
        series_path = figure_dir / f"{stem}_extended_xaxis_series.csv"
        write_series_csv(
            foreground=foreground,
            adaptive=adaptive,
            output_path=series_path,
        )
        print(f"Wrote {series_path}")

        for stacked in [False, True]:
            mode = "stacked" if stacked else "unstacked"
            png_path = figure_dir / f"{stem}_{mode}_extended_xaxis.png"
            plot_year_start, plot_year_end = write_case_png(
                trails=trails,
                foreground=foreground,
                adaptive=adaptive,
                stacked=stacked,
                path=png_path,
            )
            print(f"Wrote {png_path} (x-axis {plot_year_start}-{plot_year_end})")

        deep_vs_static = (adaptive.cumulative.iloc[-1] / static_score - 1.0) * 100.0
        deep_vs_foreground = (
            adaptive.cumulative.iloc[-1] / foreground.cumulative.iloc[-1] - 1.0
        ) * 100.0
        print(f"Deep vs static: {deep_vs_static:+.6f}%")
        print(f"Deep vs foreground-only: {deep_vs_foreground:+.6f}%")


if __name__ == "__main__":
    main()
